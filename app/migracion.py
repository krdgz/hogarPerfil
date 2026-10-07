import os
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

from dotenv import load_dotenv
import psycopg2
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
load_dotenv(BASE_DIR / ".env")

# =========================================================
# CONFIGURACION
# =========================================================
db_password = os.getenv("POSTGRES_PASSWORD", "")
if not db_password:
    raise RuntimeError(
        "Falta POSTGRES_PASSWORD en el archivo .env local. "
        "Crea un .env con la contraseña real de PostgreSQL."
    )

DB_CONFIG = {
    "host": os.getenv("POSTGRES_HOST", "localhost"),
    "port": int(os.getenv("POSTGRES_PORT", "5432")),
    "dbname": os.getenv("POSTGRES_DB", "hogares"),
    "user": os.getenv("POSTGRES_USER", "postgres"),
    "password": db_password,
}

EXCEL_PATH = "forms.xlsx"
SHEET_NAME = "LAS TUNAS"

# =========================================================
# UTILIDADES
# =========================================================
def norm(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = "".join(
        c for c in unicodedata.normalize("NFKD", s)
        if not unicodedata.combining(c)
    )
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"\s*/\s*", "/", s)
    return s


def to_bool(s):
    if s is None:
        return None
    s = norm(s)
    if s in ("si", "sí", "1", "true", "x", "yes"):
        return True
    if s in ("no", "0", "false"):
        return False
    return None


def to_int(s):
    if s is None or str(s).strip() == "":
        return None
    try:
        return int(float(str(s).replace(",", ".")))
    except Exception:
        return None


def to_decimal(s):
    if s is None or str(s).strip() == "":
        return None
    txt = re.sub(r"[^0-9.,\-]", "", str(s))
    txt = txt.replace(",", ".")
    try:
        return Decimal(txt)
    except (InvalidOperation, ValueError):
        return None


def normalizar_cedula(s):
    if s is None:
        return None
    raw = str(s).strip()
    if raw == "":
        return None
    digitos = re.sub(r"\D", "", raw)
    if digitos == "":
        return None
    if len(digitos) > 11:
        digitos = digitos[:11]
    elif len(digitos) < 11:
        digitos = digitos.ljust(11, "0")
    return digitos


def lookup(d, val):
    if val is None:
        return None
    s = norm(val)
    if s in d:
        return d[s]
    s2 = s.rstrip(":")
    if s2 in d:
        return d[s2]
    # Soporta celdas que traen el codigo numerico (2, 3, 3.0, etc.)
    try:
        code_str = str(int(float(s)))
        if code_str in d:
            return d[code_str]
    except Exception:
        pass
    return None


def safe_str(val, default=""):
    if val is None:
        return default
    s = str(val).strip()
    if s == "":
        return default
    return s


def load_catalogo(cur, table):
    cur.execute(f"SELECT id, codigo, nombre FROM {table}")
    d = {}
    for id_, codigo, nombre in cur.fetchall():
        d[norm(nombre)] = id_
        d[str(codigo).strip()] = id_
    return d


# =========================================================
# GEOGRAFIA - helpers
# =========================================================
def parse_consejo(d):
    if d is None:
        return None, None
    s = str(d).strip()
    if s == "":
        return None, None
    m = re.match(r"^(\d+)\.?\s*(.*)$", s)
    if m:
        codigo = m.group(1)
        nombre = m.group(2).strip()
    else:
        codigo = norm(s)[:10]
        nombre = s
    return codigo, nombre


# =========================================================
# MAIN
# =========================================================
def main():
    conn = psycopg2.connect(**DB_CONFIG)
    conn.autocommit = False
    cur = conn.cursor()

    # =====================================================
    # CARGAR CATALOGOS
    # =====================================================
    cat_perfil = load_catalogo(cur, "catalogo_perfil_vulnerabilidad")
    cat_color = load_catalogo(cur, "catalogo_color_piel")
    cat_parentesco = load_catalogo(cur, "catalogo_parentesco")
    cat_sexo = load_catalogo(cur, "catalogo_sexo")
    cat_escolaridad = load_catalogo(cur, "catalogo_nivel_escolaridad")
    cat_vinculacion = load_catalogo(cur, "catalogo_vinculacion_sne")
    cat_fuente = load_catalogo(cur, "catalogo_fuente_ingreso")
    cat_sector_estatal = load_catalogo(cur, "catalogo_sector_estatal")
    cat_empresa_mixta = load_catalogo(cur, "catalogo_empresa_mixta")
    cat_sector_agro = load_catalogo(cur, "catalogo_sector_agropecuario")
    cat_sector_no_estatal = load_catalogo(cur, "catalogo_sector_no_estatal")
    cat_trabajo_informal = load_catalogo(cur, "catalogo_trabajo_informal")
    cat_motivo = load_catalogo(cur, "catalogo_motivo_no_trabaja")

    # =====================================================
    # CARGAR GEOGRAFIA EXISTENTE
    # =====================================================
    prov_by_nombre = {}
    cur.execute("SELECT id, nombre FROM provincia")
    for id_, nombre in cur.fetchall():
        prov_by_nombre[norm(nombre)] = id_

    mun_by_nombre = {}
    mun_by_codigo = {}
    cur.execute("SELECT id, codigo, nombre, provincia_id FROM municipio")
    for id_, codigo, nombre, prov_id in cur.fetchall():
        mun_by_nombre[(prov_id, norm(nombre))] = id_
        mun_by_codigo[(prov_id, codigo)] = id_

    cons_by_nombre = {}
    cons_by_codigo = {}
    cur.execute("SELECT id, codigo, nombre, municipio_id FROM consejo_popular")
    for id_, codigo, nombre, muni_id in cur.fetchall():
        cons_by_nombre[(muni_id, norm(nombre))] = id_
        cons_by_codigo[(muni_id, codigo)] = id_

    circ_by_codigo = {}
    cur.execute("SELECT id, codigo, consejo_popular_id FROM circunscripcion")
    for id_, codigo, cons_id in cur.fetchall():
        circ_by_codigo[(cons_id, codigo)] = id_

    bod_by_codigo = {}
    cur.execute("SELECT id, codigo, consejo_popular_id FROM bodega")
    for id_, codigo, cons_id in cur.fetchall():
        bod_by_codigo[(cons_id, codigo)] = id_

    zona_by_nombre = {}
    cur.execute("SELECT id, nombre FROM zona_residencia")
    for id_, nombre in cur.fetchall():
        zona_by_nombre[norm(nombre)] = id_

    # =====================================================
    # FUNCIONES GET OR CREATE
    # =====================================================
    def get_provincia(nombre):
        nombre = safe_str(nombre, "DESCONOCIDA")
        n = norm(nombre)
        if n in prov_by_nombre:
            return prov_by_nombre[n]
        codigo = re.sub(r"[^A-Z0-9]", "", n.upper())[:10] or "SIN_COD"
        cur.execute("SELECT id FROM provincia WHERE codigo = %s", (codigo,))
        row = cur.fetchone()
        if row:
            id_ = row[0]
            prov_by_nombre[n] = id_
            return id_
        cur.execute(
            """
            INSERT INTO provincia (codigo, nombre)
            VALUES (%s, %s)
            RETURNING id
            """,
            (codigo, nombre),
        )
        id_ = cur.fetchone()[0]
        prov_by_nombre[n] = id_
        return id_

    def get_municipio(prov_id, nombre):
        nombre = safe_str(nombre, "SIN DATO")
        n = norm(nombre)
        key = (prov_id, n)
        if key in mun_by_nombre:
            return mun_by_nombre[key]
        codigo = re.sub(r"[^A-Z0-9]", "", n.upper())[:10] or "SIN_COD"
        cur.execute(
            "SELECT id FROM municipio WHERE provincia_id = %s AND codigo = %s",
            (prov_id, codigo),
        )
        row = cur.fetchone()
        if row:
            id_ = row[0]
            mun_by_nombre[key] = id_
            return id_
        cur.execute(
            """
            INSERT INTO municipio (codigo, nombre, provincia_id)
            VALUES (%s, %s, %s)
            RETURNING id
            """,
            (codigo, nombre, prov_id),
        )
        id_ = cur.fetchone()[0]
        mun_by_nombre[key] = id_
        mun_by_codigo[(prov_id, codigo)] = id_
        return id_

    def get_consejo(muni_id, d):
        codigo, nombre = parse_consejo(d)
        if not nombre:
            nombre = "SIN DATO"
            codigo = "SIN_COD"
        n = norm(nombre)
        key = (muni_id, n)
        if key in cons_by_nombre:
            return cons_by_nombre[key]
        cur.execute(
            "SELECT id FROM consejo_popular WHERE municipio_id = %s AND codigo = %s",
            (muni_id, codigo),
        )
        row = cur.fetchone()
        if row:
            id_ = row[0]
            cons_by_nombre[key] = id_
            return id_
        cur.execute(
            """
            INSERT INTO consejo_popular (codigo, nombre, municipio_id)
            VALUES (%s, %s, %s)
            RETURNING id
            """,
            (codigo, nombre, muni_id),
        )
        id_ = cur.fetchone()[0]
        cons_by_nombre[key] = id_
        cons_by_codigo[(muni_id, codigo)] = id_
        return id_

    def get_circunscripcion(cons_id, codigo):
        if codigo is None or str(codigo).strip() == "":
            return None
        cod = str(codigo).strip()
        key = (cons_id, cod)
        if key in circ_by_codigo:
            return circ_by_codigo[key]
        cur.execute(
            """
            INSERT INTO circunscripcion (codigo, consejo_popular_id)
            VALUES (%s, %s)
            ON CONFLICT (consejo_popular_id, codigo)
            DO UPDATE SET codigo = EXCLUDED.codigo
            RETURNING id
            """,
            (cod, cons_id),
        )
        id_ = cur.fetchone()[0]
        circ_by_codigo[key] = id_
        return id_

    def get_bodega(cons_id, codigo):
        if codigo is None or str(codigo).strip() == "":
            return None
        try:
            cod = int(float(str(codigo)))
        except Exception:
            cod = 0
        key = (cons_id, cod)
        if key in bod_by_codigo:
            return bod_by_codigo[key]
        cur.execute(
            """
            INSERT INTO bodega (codigo, consejo_popular_id)
            VALUES (%s, %s)
            ON CONFLICT (consejo_popular_id, codigo)
            DO UPDATE SET codigo = EXCLUDED.codigo
            RETURNING id
            """,
            (cod, cons_id),
        )
        id_ = cur.fetchone()[0]
        bod_by_codigo[key] = id_
        return id_

    def get_zona(nombre):
        if nombre is None or str(nombre).strip() == "":
            nombre = "Zona rural"
        n = norm(nombre)
        alias = {
            "urbana": "zona urbana",
            "rural": "zona rural",
        }
        n = alias.get(n, n)
        if n in zona_by_nombre:
            return zona_by_nombre[n]
        codigo = n.upper().replace(" ", "")[:20] or "RURAL"
        cur.execute(
            """
            INSERT INTO zona_residencia (codigo, nombre)
            VALUES (%s, %s)
            ON CONFLICT (codigo) DO UPDATE SET nombre = EXCLUDED.nombre
            RETURNING id
            """,
            (codigo, nombre),
        )
        id_ = cur.fetchone()[0]
        zona_by_nombre[n] = id_
        return id_

    # =====================================================
    # LEER EXCEL
    # =====================================================
    wb = load_workbook(EXCEL_PATH, data_only=True, read_only=True)
    ws = wb[SHEET_NAME]

    header_row = None
    headers = None
    for row in ws.iter_rows(min_row=1, max_row=15):
        vals = [cell.value for cell in row]
        if len(vals) < 2:
            continue
        if norm(vals[0]) in ("no.", "no", "nro", "nro.") and \
           "provincia" in norm(vals[1] or ""):
            header_row = row[0].row
            headers = vals
            break

    if header_row is None:
        raise RuntimeError("No se encontro la fila de encabezados en el Excel.")

    print(f"Fila de encabezados detectada: {header_row}")

    idx_map = {norm(h): i for i, h in enumerate(headers) if h is not None}

    def col(name_options, requerido=False):
        for opt in name_options:
            n = norm(opt)
            if n in idx_map:
                return idx_map[n]
        if requerido:
            print(f"[WARN] Columna no encontrada: {name_options}")
        return None

    C_NO = col(["No.", "No"])
    C_PROV = col(["Provincia"])
    C_MUNI = col(["Municipio"])
    C_CONS = col(["Consejo Popular"])
    C_CIRC = col(["Circunscripción", "Circunscripcion"])
    C_ZONA = col(["Zona de residencia"])
    C_BOD = col(["Número de bodega", "Numero de bodega"])
    C_NOMBRE = col(["Nombre y Apellidos", "Nombre y Apellidos "])
    C_CEDULA = col(["Carné de Identidad", "Carne de Identidad"])
    C_DIR = col(["Dirección particular", "Direccion particular"])
    C_PERFIL = col(["Perfil de vulnerabilidad (Seleccionar)", "Perfil de vulnerabilidad"])
    C_COLOR = col(["Color de piel"])
    C_PARENT = col(["Parentesco respecto al jefe del hogar"])
    C_SEXO = col(["Sexo"])
    C_ESCOL = col(["Nivel de escolaridad vencido"])
    C_VINC = col([
        "NNA. Vínculo con el Sistema Nacional de Enseñanza/ Cuidado de la primera",
        "NNA. Vinculo con el Sistema Nacional de Ensenanza/ Cuidado de la primera",
    ])
    C_REG = col(["Registro de consumidores"])
    C_FUENTE = col(["Fuente de los ingresos que recibe"])
    C_INGRESO = col(["Ingreso mensual que recibe en CUP"])
    C_SEC_EST = col(["Sector Estatal"])
    C_EMP_MIX = col(["Empresa mixta"])
    C_SEC_AGRO = col(["Sector agropecuario"])
    C_SEC_NO_EST = col(["Sector No estatal"])
    C_TRAB_INF = col(["Trabajos informales", "Trabajos <br>informales"])
    C_MOTIVO = col([
        "Indique el motivo principal por no trabajar (SELECCIONAR)",
        "Indique el motivo principal por no trabajar",
    ])
    C_PROCEDE = col(["procede", "Procede"])

    if C_NO is None:        C_NO = 0
    if C_PROV is None:      C_PROV = 1
    if C_MUNI is None:      C_MUNI = 2
    if C_CONS is None:      C_CONS = 3
    if C_CIRC is None:      C_CIRC = 4
    if C_ZONA is None:      C_ZONA = 5   # Columna F
    if C_BOD is None:       C_BOD = 6
    if C_NOMBRE is None:    C_NOMBRE = 7
    if C_CEDULA is None:    C_CEDULA = 8
    if C_DIR is None:       C_DIR = 9

    print("Indices -> NO:", C_NO, "PROV:", C_PROV, "MUNI:", C_MUNI,
          "CONS:", C_CONS, "CIRC:", C_CIRC, "ZONA:", C_ZONA)

    GASTO_HEADERS = [
        ("Monto en CUP", "Orden de importancia"),
        ("Monto en CUP2", "Orden de importancia3"),
        ("Monto en CUP3", "Orden de importancia4"),
        ("Monto en CUP32", "Orden de importancia44"),
        ("Monto en CUP4", "Orden de importancia6"),
        ("Monto en CUP42", "Orden de importancia64"),
        ("Monto en CUP422", "Orden de importancia643"),
        ("Monto en CUP5", "Orden de importancia7"),
    ]

    FOOD_HEADERS = [
        ("¿Le gusta?", "Lo encontró en el mercado", "Frecuencia semanal de consumo habitual (días, 0-7)"),
        ("¿Le gusta? 2", "Lo encontró en el mercado 2", "Frecuencia semanal de consumo habitual (días, 0-7) 2"),
        ("¿Le gusta? 3", "Lo encontró en el mercado 3", "Frecuencia semanal de consumo habitual (días, 0-7) 3"),
        ("¿Le gusta? 4", "Lo encontró en el mercado 4", "Frecuencia semanal de consumo habitual (días, 0-7) 4"),
        ("¿Le gusta? 5", "Lo encontró en el mercado 5", "Frecuencia semanal de consumo habitual (días, 0-7) 5"),
        ("¿Le gusta? 6", "Lo encontró en el mercado 6", "Frecuencia semanal de consumo habitual (días, 0-7) 6"),
        ("¿Le gusta? 7", "Lo encontró en el mercado 7", "Frecuencia semanal de consumo habitual (días, 0-7) 7"),
        ("¿Le gusta? 72", "Lo encontró en el mercado 73", "Frecuencia semanal de consumo habitual (días, 0-7) 74"),
        ("¿Le gusta? 8", "Lo encontró en el mercado 8", "Frecuencia semanal de consumo habitual (días, 0-7) 8"),
    ]

    ESTRATEGIAS_HEADERS = [
        "Recurrir a alimentos menos preferidos y más baratos",
        "Pedir alimentos prestados o ayuda de familiares/amigos",
        "Reducir el número de comidas diarias",
        "Disminuir el tamaño de las porciones",
        "Adultos restringen comida para priorizar niños, ancianos o enfermos",
    ]

    # =====================================================
    # LEER FILAS DE DATOS
    # =====================================================
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or all(v is None for v in row):
            continue

        def get(idx):
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        no = get(C_NO)
        nombre = get(C_NOMBRE)
        if no is None and nombre is None:
            continue
        if no is None:
            continue

        rec = {
            "no": no,
            "provincia": get(C_PROV),
            "municipio": get(C_MUNI),
            "consejo": get(C_CONS),
            "circunscripcion": get(C_CIRC),
            "zona": get(C_ZONA),
            "bodega": get(C_BOD),
            "nombre": nombre,
            "cedula": get(C_CEDULA),
            "direccion": get(C_DIR),
            "perfil": get(C_PERFIL),
            "color": get(C_COLOR),
            "parentesco": get(C_PARENT),
            "sexo": get(C_SEXO),
            "escolaridad": get(C_ESCOL),
            "vinculacion": get(C_VINC),
            "registro": get(C_REG),
            "fuente": get(C_FUENTE),
            "ingreso": get(C_INGRESO),
            "sector_estatal": get(C_SEC_EST),
            "empresa_mixta": get(C_EMP_MIX),
            "sector_agro": get(C_SEC_AGRO),
            "sector_no_estatal": get(C_SEC_NO_EST),
            "trabajo_informal": get(C_TRAB_INF),
            "motivo_no_trabaja": get(C_MOTIVO),
            "procede": get(C_PROCEDE),
            "gastos": [],
            "alimentos": [],
            "estrategias": [],
        }

        for i, (mh, oh) in enumerate(GASTO_HEADERS):
            m_idx = idx_map.get(norm(mh))
            o_idx = idx_map.get(norm(oh))
            monto = row[m_idx] if m_idx is not None and m_idx < len(row) else None
            orden = row[o_idx] if o_idx is not None and o_idx < len(row) else None
            rec["gastos"].append((i + 1, monto, orden))

        for i, (gh, eh, fh) in enumerate(FOOD_HEADERS):
            g_idx = idx_map.get(norm(gh))
            e_idx = idx_map.get(norm(eh))
            f_idx = idx_map.get(norm(fh))
            gusta = row[g_idx] if g_idx is not None and g_idx < len(row) else None
            encontrado = row[e_idx] if e_idx is not None and e_idx < len(row) else None
            frec = row[f_idx] if f_idx is not None and f_idx < len(row) else None
            rec["alimentos"].append((i + 1, gusta, encontrado, frec))

        for i, h in enumerate(ESTRATEGIAS_HEADERS):
            idx = idx_map.get(norm(h))
            val = row[idx] if idx is not None and idx < len(row) else None
            rec["estrategias"].append((i + 1, val))

        rows.append(rec)

    print(f"Filas de datos leidas: {len(rows)}")

    # =====================================================
    # AGRUPAR POR HOGAR (No.)
    # =====================================================
    hogares = {}
    for rec in rows:
        no = str(rec["no"]).strip()
        if no not in hogares:
            hogares[no] = []
        hogares[no].append(rec)

    print(f"Hogares detectados: {len(hogares)}")

    # =====================================================
    # INSERTAR / ACTUALIZAR
    # =====================================================
    for no, personas in hogares.items():
        first = personas[0]

        prov_id = get_provincia(first["provincia"])
        muni_id = get_municipio(prov_id, first["municipio"])
        cons_id = get_consejo(muni_id, first["consejo"])
        circ_id = get_circunscripcion(cons_id, first["circunscripcion"])
        bod_id = get_bodega(cons_id, first["bodega"])
        zona_id = get_zona(first["zona"])
        direccion = first["direccion"]
        procede = to_bool(first["procede"])

        cur.execute(
            """
            INSERT INTO hogar_nucleo (
                codigo, direccion, provincia_id, municipio_id,
                consejo_popular_id, circunscripcion_id, bodega_id,
                zona_residencia_id, procede_ayuda
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (codigo) DO UPDATE SET
                direccion = EXCLUDED.direccion,
                provincia_id = EXCLUDED.provincia_id,
                municipio_id = EXCLUDED.municipio_id,
                consejo_popular_id = EXCLUDED.consejo_popular_id,
                circunscripcion_id = EXCLUDED.circunscripcion_id,
                bodega_id = EXCLUDED.bodega_id,
                zona_residencia_id = EXCLUDED.zona_residencia_id,
                procede_ayuda = EXCLUDED.procede_ayuda
            RETURNING id
            """,
            (no, direccion, prov_id, muni_id, cons_id, circ_id, bod_id, zona_id, procede),
        )
        hogar_id = cur.fetchone()[0]

        cur.execute("DELETE FROM gasto_hogar WHERE hogar_id = %s", (hogar_id,))
        cur.execute("DELETE FROM diversidad_alimentaria_hogar WHERE hogar_id = %s", (hogar_id,))
        cur.execute("DELETE FROM estrategia_afrontamiento_hogar WHERE hogar_id = %s", (hogar_id,))

        for tipo_id, monto, orden in first["gastos"]:
            if monto is None or str(monto).strip() == "":
                continue
            m = to_decimal(monto)
            if m is None:
                continue
            o = to_int(orden)
            if o is None or o < 1 or o > 3:
                o = 1
            cur.execute(
                """
                INSERT INTO gasto_hogar (hogar_id, tipo_gasto_id, monto_cup, orden_importancia)
                VALUES (%s, %s, %s, %s)
                ON CONFLICT (hogar_id, tipo_gasto_id) DO NOTHING
                """,
                (hogar_id, tipo_id, m, o),
            )

        for grupo_id, gusta, encontrado, frec in first["alimentos"]:
            if gusta is None and encontrado is None and (frec is None or str(frec).strip() == ""):
                continue
            f = to_int(frec)
            if f is None:
                f = 0
            f = max(0, min(7, f))
            cur.execute(
                """
                INSERT INTO diversidad_alimentaria_hogar (
                    hogar_id, grupo_alimento_id, gusta,
                    encontrado_en_mercado, frecuencia_semanal_dias
                )
                VALUES (%s, %s, %s, %s, %s)
                ON CONFLICT (hogar_id, grupo_alimento_id) DO NOTHING
                """,
                (hogar_id, grupo_id, to_bool(gusta), to_bool(encontrado), f),
            )

        for est_id, val in first["estrategias"]:
            if val is None or str(val).strip() == "":
                continue
            cur.execute(
                """
                INSERT INTO estrategia_afrontamiento_hogar (hogar_id, estrategia_id, aplica)
                VALUES (%s, %s, %s)
                ON CONFLICT (hogar_id, estrategia_id) DO NOTHING
                """,
                (hogar_id, est_id, to_bool(val) or False),
            )

        # Personas
        cedulas_vistas = set()
        for p in personas:
            cedula = normalizar_cedula(p["cedula"])
            if cedula is not None and cedula in cedulas_vistas:
                print(f"[WARN] Hogar {no}: cedula duplicada tras normalizar -> {cedula}. Se omite.")
                continue
            if cedula is not None:
                cedulas_vistas.add(cedula)

            nombre = p["nombre"]
            color_id = lookup(cat_color, p["color"])
            parentesco_id = lookup(cat_parentesco, p["parentesco"])
            sexo_id = lookup(cat_sexo, p["sexo"])
            escolaridad_id = lookup(cat_escolaridad, p["escolaridad"])
            vinculacion_id = lookup(cat_vinculacion, p["vinculacion"])
            fuente_id = lookup(cat_fuente, p["fuente"])
            ingreso = to_decimal(p["ingreso"])
            registro = to_bool(p["registro"])
            motivo_id = lookup(cat_motivo, p["motivo_no_trabaja"])

            ocupaciones = []

            if p["sector_estatal"] not in (None, "", "No", "NO"):
                cat_id = lookup(cat_sector_estatal, p["sector_estatal"])
                if cat_id:
                    ocupaciones.append(("sector_estatal", "sector_estatal_id", cat_id))

            if to_bool(p["empresa_mixta"]) is True:
                cat_id = cat_empresa_mixta.get("si")
                if cat_id:
                    ocupaciones.append(("empresa_mixta", "empresa_mixta_id", cat_id))

            if p["sector_agro"] not in (None, "", "No", "NO"):
                cat_id = lookup(cat_sector_agro, p["sector_agro"])
                if cat_id:
                    ocupaciones.append(("sector_agropecuario", "sector_agropecuario_id", cat_id))

            if p["sector_no_estatal"] not in (None, "", "No", "NO"):
                cat_id = lookup(cat_sector_no_estatal, p["sector_no_estatal"])
                if cat_id:
                    ocupaciones.append(("sector_no_estatal", "sector_no_estatal_id", cat_id))

            if to_bool(p["trabajo_informal"]) is True:
                cat_id = cat_trabajo_informal.get("si")
                if cat_id:
                    ocupaciones.append(("trabajo_informal", "trabajo_informal_id", cat_id))

            trabaja = True if ocupaciones else (False if motivo_id else None)

            cur.execute(
                """
                INSERT INTO persona (
                    hogar_id, nombre_apellidos, cedula, color_piel_id,
                    parentesco_jefe_id, sexo_id, nivel_escolaridad_id,
                    vinculacion_sne_id, esta_en_registro_consumidor,
                    fuente_ingreso_id, ingreso_mensual_cup, trabaja,
                    motivo_no_trabaja_id
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (hogar_id, cedula) DO UPDATE SET
                    nombre_apellidos = EXCLUDED.nombre_apellidos,
                    color_piel_id = EXCLUDED.color_piel_id,
                    parentesco_jefe_id = EXCLUDED.parentesco_jefe_id,
                    sexo_id = EXCLUDED.sexo_id,
                    nivel_escolaridad_id = EXCLUDED.nivel_escolaridad_id,
                    vinculacion_sne_id = EXCLUDED.vinculacion_sne_id,
                    esta_en_registro_consumidor = EXCLUDED.esta_en_registro_consumidor,
                    fuente_ingreso_id = EXCLUDED.fuente_ingreso_id,
                    ingreso_mensual_cup = EXCLUDED.ingreso_mensual_cup,
                    trabaja = EXCLUDED.trabaja,
                    motivo_no_trabaja_id = EXCLUDED.motivo_no_trabaja_id
                RETURNING id
                """,
                (
                    hogar_id, nombre, cedula, color_id,
                    parentesco_id, sexo_id, escolaridad_id,
                    vinculacion_id, registro, fuente_id,
                    ingreso, trabaja, motivo_id,
                ),
            )
            persona_id = cur.fetchone()[0]

            cur.execute("DELETE FROM persona_vulnerabilidad WHERE persona_id = %s", (persona_id,))
            cur.execute("DELETE FROM persona_ocupacion WHERE persona_id = %s", (persona_id,))

            if p["perfil"] is not None and str(p["perfil"]).strip() != "":
                perfil_texto = str(p["perfil"]).strip().rstrip(":")
                perfil_completo_id = lookup(cat_perfil, perfil_texto)
                if perfil_completo_id is not None:
                    perfiles = [(perfil_texto, perfil_completo_id)]
                else:
                    perfiles = [
                        (perfil, lookup(cat_perfil, perfil.strip().rstrip(":")))
                        for perfil in re.split(r"[;,]", perfil_texto)
                    ]

                for perfil, cat_id in perfiles:
                    perfil = perfil.strip().rstrip(":")
                    if not perfil:
                        continue
                    if cat_id is not None:
                        cur.execute(
                            """
                            INSERT INTO persona_vulnerabilidad (persona_id, perfil_vulnerabilidad_id)
                            VALUES (%s, %s)
                            ON CONFLICT DO NOTHING
                            """,
                            (persona_id, cat_id),
                        )
                    else:
                        print(f"[WARN] Hogar {no}, persona {cedula}: perfil no reconocido -> {perfil!r}")

            for idx, (tipo, col_name, cat_id) in enumerate(ocupaciones):
                es_principal = idx == 0
                cur.execute(
                    f"""
                    INSERT INTO persona_ocupacion (
                        persona_id, es_principal, tipo_ocupacion, {col_name}
                    )
                    VALUES (%s, %s, %s, %s)
                    """,
                    (persona_id, es_principal, tipo, cat_id),
                )

        conn.commit()
        print(f"Hogar {no} procesado: {len(personas)} personas")

    cur.close()
    conn.close()
    print("Migracion completada.")


if __name__ == "__main__":
    main()